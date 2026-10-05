"""Читання й запис книги «Лічильники по підстанціям 2026.xlsm».

Читання — через openpyxl (Excel не потрібен, книга може бути відкрита у когось іншого).
Запис — через Excel COM (pywin32): лише так безпечно зберігаються макроси VBA, кнопки,
об'єднані комірки й форматування. Кожна операція відкриває книгу в окремому невидимому
екземплярі Excel, змінює, зберігає й закриває його.
"""

from __future__ import annotations

import logging
import os
import shutil
import tempfile
from datetime import date, datetime
from pathlib import Path
from typing import Dict, Iterator, List, Optional, Sequence, Tuple

from meters_model import (
    ACT_HEADERS,
    ActEntry,
    SH_ACTS,
    act_entry_for_record,
    act_fields_changed,
    COL_NUM,
    COL_SUBST,
    FIRST_DATA_ROW,
    HDR_BOTTOM,
    LAST_COL,
    REF_FIRST_ROW,
    REF_KEY_ROW,
    SH_BASE,
    SH_HISTORY,
    SH_REF,
    Snapshot,
    caption_of,
    format_value,
    parse_cell_text,
)

WORKBOOK_GLOB = "Лічильники по підстанціям*.xlsm"
HISTORY_HEADERS = ["Дата Час", "Користувач", "Рядок бази", "Поле", "Старе значення", "Нове значення"]

# Константи Excel
_XL_DOWN = -4121
_XL_FORMULAS = -4123
_XL_PART = 2
_XL_BY_ROWS = 1
_XL_PREVIOUS = 2
_MSO_AUTOMATION_SECURITY_FORCE_DISABLE = 3


class MetersError(Exception):
    """Помилка роботи з книгою (повідомлення призначене для користувача)."""


def find_workbook(app_dir: Path) -> Optional[Path]:
    found = sorted(app_dir.glob(WORKBOOK_GLOB))
    return found[0] if found else None


# ---------------------------------------------------------------- читання
def _clean(v: object) -> str:
    return format_value(v)


def load_snapshot(path: Path) -> Snapshot:
    try:
        import openpyxl  # noqa: WPS433
    except ImportError as e:  # pragma: no cover
        raise MetersError("Не встановлено бібліотеку openpyxl.") from e

    if not path.is_file():
        raise MetersError("Файл не знайдено: {}".format(path))

    tmp_dir: Optional[str] = None
    src = path
    try:
        try:
            wb = openpyxl.load_workbook(str(src), data_only=True)
        except PermissionError:
            # книга відкрита в Excel з блокуванням — читаємо копію
            tmp_dir = tempfile.mkdtemp(prefix="meters_")
            src = Path(tmp_dir) / path.name
            shutil.copyfile(str(path), str(src))
            wb = openpyxl.load_workbook(str(src), data_only=True)
        except Exception as e:
            raise MetersError("Не вдалося відкрити книгу: {}".format(e)) from e

        try:
            if SH_BASE not in wb.sheetnames:
                raise MetersError('У книзі немає аркуша «{}».'.format(SH_BASE))
            ws = wb[SH_BASE]
            snap = _read_base(ws)
            if SH_REF in wb.sheetnames:
                snap.ref_lists = _read_ref(wb[SH_REF])
        finally:
            wb.close()
        return snap
    finally:
        if tmp_dir:
            shutil.rmtree(tmp_dir, ignore_errors=True)


def _read_base(ws) -> Snapshot:  # type: ignore[no-untyped-def]
    max_r = max(ws.max_row, FIRST_DATA_ROW - 1)
    grid: List[List[object]] = []
    for row in ws.iter_rows(min_row=1, max_row=max_r, max_col=LAST_COL, values_only=True):
        r = list(row)
        r.extend([None] * (LAST_COL - len(r)))
        grid.append(r)
    while len(grid) < max_r:
        grid.append([None] * LAST_COL)

    row1_wide = [False] * LAST_COL
    merge_anchor: Dict[Tuple[int, int], int] = {}
    for rng in ws.merged_cells.ranges:
        c1, c2 = rng.min_col, min(rng.max_col, LAST_COL)
        r1, r2 = rng.min_row, rng.max_row
        if c1 > LAST_COL:
            continue
        anchor = grid[r1 - 1][c1 - 1] if r1 - 1 < len(grid) else None
        if r1 == 1 and rng.max_col > rng.min_col:
            for c in range(c1, c2 + 1):
                row1_wide[c - 1] = True
        for r in range(r1, min(r2, len(grid)) + 1):
            for c in range(c1, c2 + 1):
                if r != r1 or c != c1:
                    grid[r - 1][c - 1] = anchor
                if r1 >= FIRST_DATA_ROW and r > r1:
                    merge_anchor[(r, c)] = r1

    hdr = [[_clean(v) for v in grid[i]] for i in range(HDR_BOTTOM)]

    last = FIRST_DATA_ROW - 1
    for r in range(len(grid), FIRST_DATA_ROW - 1, -1):
        if any(v is not None and str(v).strip() != "" for v in grid[r - 1]):
            last = r
            break
    values = [
        [_clean(v) for v in grid[r - 1]] for r in range(FIRST_DATA_ROW, last + 1)
    ]
    return Snapshot(hdr=hdr, row1_wide=row1_wide, values=values, merge_anchor=merge_anchor)


def _read_ref(ws) -> Dict[str, List[str]]:  # type: ignore[no-untyped-def]
    out: Dict[str, List[str]] = {}
    max_c = ws.max_column
    max_r = ws.max_row
    for c in range(1, max_c + 1):
        key = _clean(ws.cell(row=REF_KEY_ROW, column=c).value).strip()
        if not key or key in out:
            continue
        seen: Dict[str, None] = {}
        for r in range(REF_FIRST_ROW, max_r + 1):
            s = _clean(ws.cell(row=r, column=c).value).strip()
            if s:
                seen.setdefault(s, None)
        if seen:
            out[key] = list(seen)
    return out


# ---------------------------------------------------------------- запис (Excel COM)
class _Session:
    """Відкрита книга в окремому невидимому екземплярі Excel."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.app = None
        self.wb = None

    def __enter__(self) -> "_Session":
        try:
            import pythoncom  # type: ignore
            import win32com.client  # type: ignore
        except ImportError as e:
            raise MetersError(
                "Для запису в Excel потрібна бібліотека pywin32 (pip install pywin32)."
            ) from e
        pythoncom.CoInitialize()
        try:
            self.app = win32com.client.DispatchEx("Excel.Application")
        except Exception as e:
            pythoncom.CoUninitialize()
            raise MetersError("Не вдалося запустити Microsoft Excel: {}".format(e)) from e
        try:
            self.app.Visible = False
            self.app.DisplayAlerts = False
            self.app.EnableEvents = False
            self.app.AutomationSecurity = _MSO_AUTOMATION_SECURITY_FORCE_DISABLE
            self.wb = self.app.Workbooks.Open(str(self.path), 0, False)
            if self.wb.ReadOnly:
                raise MetersError(
                    "Книга відкрита в Excel іншим користувачем або програмою.\n"
                    "Закрийте її й повторіть дію."
                )
        except MetersError:
            self._quit()
            raise
        except Exception as e:
            self._quit()
            raise MetersError("Не вдалося відкрити книгу для запису: {}".format(e)) from e
        return self

    def __exit__(self, exc_type, exc, tb) -> None:  # type: ignore[no-untyped-def]
        try:
            if exc_type is None and self.wb is not None:
                self.wb.Save()
        except Exception as e:
            self._quit()
            raise MetersError("Не вдалося зберегти книгу: {}".format(e)) from e
        self._quit()

    def _quit(self) -> None:
        try:
            import pythoncom  # type: ignore

            if self.wb is not None:
                try:
                    self.wb.Close(False)
                except Exception:
                    logging.exception("Excel: Close")
            if self.app is not None:
                try:
                    self.app.Quit()
                except Exception:
                    logging.exception("Excel: Quit")
            self.wb = None
            self.app = None
            pythoncom.CoUninitialize()
        except Exception:
            logging.exception("Excel: cleanup")

    @property
    def base(self):  # type: ignore[no-untyped-def]
        return self.wb.Worksheets(SH_BASE)

    def last_data_row(self) -> int:
        ws = self.base
        f = ws.Cells.Find(
            "*", ws.Cells(1, 1), _XL_FORMULAS, _XL_PART, _XL_BY_ROWS, _XL_PREVIOUS
        )
        last = f.Row if f is not None else FIRST_DATA_ROW - 1
        return max(last, FIRST_DATA_ROW - 1)

    def merged_cell(self, r: int, c: int):  # type: ignore[no-untyped-def]
        rng = self.base.Cells(r, c)
        if rng.MergeCells:
            rng = rng.MergeArea.Cells(1, 1)
        return rng

    def text_to_cell(self, r: int, c: int, s: str) -> None:
        rng = self.merged_cell(r, c)
        value, force_text = parse_cell_text(s)
        if value is None:
            rng.ClearContents()
        elif force_text or rng.NumberFormat == "@":
            rng.NumberFormat = "@"
            rng.Value = (s or "").strip()
        elif isinstance(value, date) and not isinstance(value, datetime):
            rng.Value = datetime(value.year, value.month, value.day)
        else:
            rng.Value = value

    def cell_text(self, r: int, c: int) -> str:
        return format_value(self.merged_cell(r, c).Value)

    def renumber(self) -> None:
        ws = self.base
        last = self.last_data_row()
        for i, r in enumerate(range(FIRST_DATA_ROW, last + 1), start=1):
            cell = ws.Cells(r, COL_NUM)
            if cell.Value != i:
                cell.Value = i

    def _append(
        self,
        sheet: str,
        headers: Sequence[str],
        rows: Sequence[Sequence[object]],
        text_cols: Sequence[int] = (),
    ) -> None:
        """Дописує рядки в кінець службового аркуша (створює його із заголовками)."""
        if not rows:
            return
        try:
            ws = self.wb.Worksheets(sheet)
        except Exception:
            ws = self.wb.Worksheets.Add(After=self.wb.Worksheets(self.wb.Worksheets.Count))
            ws.Name = sheet
            for i, h in enumerate(headers, start=1):
                ws.Cells(1, i).Value = h
            ws.Rows(1).Font.Bold = True
            ws.Cells(1, 1).Select()
        first = ws.Cells(ws.Rows.Count, 1).End(-4162).Row + 1  # xlUp
        for k, row in enumerate(rows):
            r = first + k
            for c, v in enumerate(row, start=1):
                cell = ws.Cells(r, c)
                if c == 1:
                    cell.NumberFormat = "dd.mm.yyyy hh:mm:ss"
                elif c in text_cols:
                    cell.NumberFormat = "@"
                cell.Value = v

    def log(self, rows: Sequence[Sequence[object]]) -> None:
        self._append(SH_HISTORY, HISTORY_HEADERS, rows, text_cols=(5, 6))

    def log_acts(self, entries: Sequence[ActEntry]) -> None:
        # усі стовпці, крім дати-часу, — текст (номера актів і EIC не повинні губити нулі)
        self._append(
            SH_ACTS,
            ACT_HEADERS,
            [e.as_row() for e in entries],
            text_cols=tuple(range(2, len(ACT_HEADERS) + 1)),
        )


def _user() -> str:
    return os.environ.get("USERNAME") or os.environ.get("USER") or ""


def save_record(
    path: Path, snap: Snapshot, r: int, changes: Dict[int, str]
) -> int:
    """Записує змінені поля рядка r; повертає кількість записаних полів."""
    if not changes:
        return 0
    now = datetime.now()
    hist: List[List[object]] = []
    with _Session(path) as s:
        for c, new in sorted(changes.items()):
            old = s.cell_text(r, c)
            s.text_to_cell(r, c, new)
            hist.append([now, _user(), r, caption_of(snap, c), old, new])
        s.log(hist)
        if act_fields_changed(changes) and changes.get(57, snap.text(r, 57)).strip():
            # змінено дані акту (номер, дата, виконавець) — це подія в журналі актів
            s.log_acts([act_entry_for_record(snap, r, changes, user=_user(), when=now)])
    return len(changes)


def register_act(path: Path, entry: ActEntry) -> None:
    """Дописує подію в журнал «Акти» (наприклад, після перейменування скану)."""
    with _Session(path) as s:
        s.log_acts([entry])


def insert_record(path: Path, subst: str) -> int:
    """Додає порожній рядок у кінець блоку підстанції (або в кінець таблиці)."""
    with _Session(path) as s:
        ws = s.base
        gl = 0
        last = s.last_data_row()
        for r in range(FIRST_DATA_ROW, last + 1):
            if str(s.merged_cell(r, COL_SUBST).Value or "").strip() == subst:
                gl = r
        if subst and gl >= FIRST_DATA_ROW:
            new_row = gl + 1
            ws.Rows(new_row).Insert(_XL_DOWN, 0)
            ws.Rows(new_row).ClearContents()
            cell = ws.Cells(gl, COL_SUBST)
            if cell.MergeCells:
                area = cell.MergeArea
                top = area.Row
                value = area.Cells(1, 1).Value
                area.UnMerge()
                ws.Range(ws.Cells(top, COL_SUBST), ws.Cells(new_row, COL_SUBST)).Merge()
                ws.Cells(top, COL_SUBST).Value = value
            else:
                ws.Cells(new_row, COL_SUBST).Value = subst
        else:
            new_row = last + 1
            if subst:
                ws.Cells(new_row, COL_SUBST).Value = subst
        s.renumber()
        s.log([[datetime.now(), _user(), new_row, "(новий рядок)", "", subst]])
    return new_row


def delete_record(path: Path, r: int) -> None:
    with _Session(path) as s:
        ws = s.base
        old = "; ".join(
            x for x in (s.cell_text(r, 3), s.cell_text(r, 5), s.cell_text(r, 7)) if x
        )
        cell = ws.Cells(r, COL_SUBST)
        carry = None
        if cell.MergeCells:
            area = cell.MergeArea
            if area.Row == r and area.Rows.Count > 1:
                carry = area.Cells(1, 1).Value  # значення якоря пропадає разом із рядком
        ws.Rows(r).Delete()
        if carry is not None:
            ws.Cells(r, COL_SUBST).Value = carry
        s.renumber()
        s.log([[datetime.now(), _user(), r, "(видалено рядок)", old, ""]])


def iter_sheet(path: Path, sheet: str, ncols: int) -> Iterator[Tuple[str, ...]]:
    """Рядки службового аркуша (без заголовка) — для перегляду у вікні."""
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    try:
        if sheet not in wb.sheetnames:
            return
        for i, row in enumerate(wb[sheet].iter_rows(max_col=ncols, values_only=True)):
            if i == 0:
                continue
            yield tuple(
                v.strftime("%d.%m.%Y %H:%M:%S") if isinstance(v, datetime) else format_value(v)
                for v in row
            )
    finally:
        wb.close()


def iter_history(path: Path) -> Iterator[Tuple[str, ...]]:
    return iter_sheet(path, SH_HISTORY, len(HISTORY_HEADERS))


def iter_acts(path: Path) -> Iterator[Tuple[str, ...]]:
    return iter_sheet(path, SH_ACTS, len(ACT_HEADERS))


# ---------------------------------------------------------------- друк у PDF
def export_cards_pdf(cards: Sequence[Sequence[Tuple[str, Sequence[Tuple[str, str]]]]], pdf_path: Path) -> None:
    """Зберігає картки записів у PDF через Excel (тимчасова книга; дані не змінюються)."""
    if not cards:
        raise MetersError("Немає записів для друку.")
    try:
        import pythoncom  # type: ignore
        import win32com.client  # type: ignore
    except ImportError as e:
        raise MetersError("Для друку потрібна бібліотека pywin32.") from e

    pythoncom.CoInitialize()
    app = None
    wb = None
    try:
        try:
            app = win32com.client.DispatchEx("Excel.Application")
        except Exception as e:
            raise MetersError("Не вдалося запустити Microsoft Excel: {}".format(e)) from e
        app.Visible = False
        app.DisplayAlerts = False
        wb = app.Workbooks.Add()
        ws = wb.Worksheets(1)

        rows: List[List[str]] = []
        titles: List[int] = []
        headers: List[int] = []
        breaks: List[int] = []
        for n, card in enumerate(cards):
            if n:
                breaks.append(len(rows) + 1)
            lookup = {lab: val for _g, items in card for lab, val in items}
            title = " — ".join(
                x for x in (
                    "Дані по лічильнику",
                    lookup.get("Назва підстанції", ""),
                    lookup.get("Назва приєднання", ""),
                ) if x
            )
            rows.append([title, "", "", ""])
            titles.append(len(rows))
            for gname, items in card:
                rows.append([gname, "", "", ""])
                headers.append(len(rows))
                for i in range(0, len(items), 2):
                    a = items[i]
                    b = items[i + 1] if i + 1 < len(items) else ("", "")
                    rows.append([a[0], a[1], b[0], b[1]])
            rows.append(["", "", "", ""])

        rng = ws.Range(ws.Cells(1, 1), ws.Cells(len(rows), 4))
        rng.NumberFormat = "@"
        rng.Value = tuple(tuple(r) for r in rows)
        rng.Font.Size = 9
        rng.VerticalAlignment = -4108  # xlCenter
        for col, w in zip((1, 2, 3, 4), (32, 24, 32, 24)):
            ws.Columns(col).ColumnWidth = w
        for r in titles:
            tr = ws.Range(ws.Cells(r, 1), ws.Cells(r, 4))
            tr.Merge()
            tr.Font.Size = 13
            tr.Font.Bold = True
        for r in headers:
            hr = ws.Range(ws.Cells(r, 1), ws.Cells(r, 4))
            hr.Interior.Color = 0xF7EBDD
            hr.Font.Bold = True
        ps = ws.PageSetup
        ps.Orientation = 1
        ps.Zoom = False
        ps.FitToPagesWide = 1
        ps.FitToPagesTall = False
        for r in breaks:
            ws.HPageBreaks.Add(ws.Cells(r, 1))
        pdf_path.parent.mkdir(parents=True, exist_ok=True)
        wb.ExportAsFixedFormat(0, str(pdf_path))
    except MetersError:
        raise
    except Exception as e:
        raise MetersError("Не вдалося створити PDF: {}".format(e)) from e
    finally:
        try:
            if wb is not None:
                wb.Close(False)
            if app is not None:
                app.Quit()
        except Exception:
            logging.exception("Excel: cleanup (pdf)")
        pythoncom.CoUninitialize()
