"""Читання книги «Лічильники по підстанціям» (без Excel): шапка, об'єднання, довідники."""

from datetime import datetime

import openpyxl
import pytest

import meters_excel as mx
import meters_model as mm


@pytest.fixture
def book(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = mm.SH_BASE
    ws["A1"], ws["B1"], ws["C1"] = "№", "EIC-код точки", "Назва підстанції"
    ws["F1"] = "Лічильник"
    ws.merge_cells("F1:G1")
    ws["F2"], ws["G2"] = "Тип", "№ лічильника"
    ws.merge_cells("A1:A3")
    ws["A4"], ws["B4"], ws["C4"] = 1, "62Z100938201573H", "ЕЧЕ-1"
    ws["A5"], ws["B5"] = 2, "62Z231508073446X"
    ws["A6"], ws["C6"] = 3, "ЕЧЕ-2"
    ws.merge_cells("C4:C5")
    ws["G4"] = 84760419
    ws["H4"] = 5.5
    ws["I4"] = datetime(2024, 8, 14)
    ref = wb.create_sheet(mm.SH_REF)
    ref["A4"], ref["A5"] = "Квартал", "Назва"
    for i, v in enumerate(["I", "II", "II", ""], start=6):
        ref.cell(row=i, column=1, value=v or None)
    p = tmp_path / "Лічильники по підстанціям 2026.xlsm"
    wb.save(str(p).replace(".xlsm", ".xlsx"))
    (tmp_path / "Лічильники по підстанціям 2026.xlsx").rename(p)
    return p


def test_load_snapshot(book):
    s = mx.load_snapshot(book)
    assert s.last_row == 6
    assert s.text(4, 3) == "ЕЧЕ-1" and s.text(5, 3) == "ЕЧЕ-1"  # об'єднана комірка
    assert s.is_locked(5, 3) and s.is_locked(4, 3)
    assert s.merge_anchor == {(5, 3): 4}
    assert s.row1_wide[5] and s.row1_wide[6] and not s.row1_wide[1]
    assert mm.group_of(s, 7) == "Лічильник"
    assert mm.caption_of(s, 7) == "№ лічильника"
    assert s.text(4, 7) == "84760419"
    assert s.text(4, 8) == "5,5"
    assert s.text(4, 9) == "14.08.2024"
    assert s.ref_lists == {"Квартал": ["I", "II"]}
    assert s.substations() == ["ЕЧЕ-1", "ЕЧЕ-2"]


def test_load_snapshot_errors(tmp_path):
    with pytest.raises(mx.MetersError):
        mx.load_snapshot(tmp_path / "немає.xlsm")


def test_find_workbook(book, tmp_path):
    assert mx.find_workbook(tmp_path) == book
    assert mx.find_workbook(tmp_path / "порожньо") is None
